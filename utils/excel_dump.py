import re
import os
from io import BytesIO
from pathlib import Path
import dotenv
from typing import Optional, Sequence

import aiosqlite
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

_TABLE_NAME_RE = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")

dotenv.load_dotenv()


def _quote_sqlite_ident(name: str) -> str:
    """
    Кавычки для идентификаторов SQLite.
    Для простоты/безопасности разрешаем только [A-Za-z0-9_], не давая инъекций через имя таблицы.
    """
    if not _TABLE_NAME_RE.fullmatch(name):
        raise ValueError(
            f"Недопустимое имя таблицы: {name!r}. Разрешены только буквы/цифры/подчёркивание, "
            f"и имя не должно начинаться с цифры."
        )
    return f'"{name}"'


async def ___dump_sqlite_table_to_xlsx_bytes(
        db_path: str | Path,
        table_name: str,
        *,
        sheet_name: Optional[str] = None,
        fetch_size: int = 2000,
        freeze_header: bool = True,
        auto_filter: bool = True,
        set_column_widths: bool = False,
) -> BytesIO:
    """
    Дамп всей таблицы SQLite в XLSX и возврат результата как BytesIO (без сохранения на диск).

    :return: BytesIO, указатель установлен на начало (seek(0))
    """
    db_path = Path(db_path)

    quoted_table = _quote_sqlite_ident(table_name)

    wb = Workbook(write_only=False)
    ws = wb.active
    ws.title = (sheet_name or table_name)[:31]

    async with aiosqlite.connect(db_path) as db:
        db.row_factory = aiosqlite.Row

        # Колонки
        async with db.execute(f"PRAGMA table_info({quoted_table})") as cur:
            cols_info = await cur.fetchall()

        if not cols_info:
            raise ValueError(f"Таблица {table_name!r} не найдена или не содержит колонок.")

        columns: Sequence[str] = [row["name"] for row in cols_info]

        # Заголовок
        ws.append(list(columns))
        if freeze_header:
            ws.freeze_panes = "A2"

        col_widths = [len(c) for c in columns]

        # Данные
        async with db.execute(f"SELECT * FROM {quoted_table}") as cur:
            while True:
                rows = await cur.fetchmany(fetch_size)
                if not rows:
                    break

                for r in rows:
                    values = [r[c] for c in columns]

                    # bytes -> hex-строка, чтобы Excel не ругался
                    for i, v in enumerate(values):
                        if isinstance(v, (bytes, bytearray, memoryview)):
                            values[i] = bytes(v).hex()

                    ws.append(values)

                    if set_column_widths:
                        for i, v in enumerate(values):
                            s = "" if v is None else str(v)
                            if len(s) > col_widths[i]:
                                col_widths[i] = min(len(s), 80)

    if auto_filter:
        last_col = get_column_letter(len(columns))
        ws.auto_filter.ref = f"A1:{last_col}1"

    if set_column_widths:
        for i, w in enumerate(col_widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = max(8, min(w + 2, 80))

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


async def xlsx_dump():
    return await ___dump_sqlite_table_to_xlsx_bytes(
        os.getenv("DATABASE_NAME"),
        os.getenv("DATABASE_USERS_TABLE")
    )
